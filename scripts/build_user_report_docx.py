"""将用户版 Markdown 实验报告转换为符合论文正文格式的 Word 文档。

格式依据用户给出的截图要求：
- 中文正文：宋体，小四号，首行缩进 2 字符，1.5 倍行距；
- 英文正文：Times New Roman，小四号；
- 中文一级标题：黑体四号顶格；
- 中文二级标题：黑体小四号顶格；
- 三级及以下标题：宋体小四号加粗顶格；
- 表格和代码块在上述基础上做适合 Word 阅读的排版。
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = PROJECT_ROOT / "reports" / "user_experiment_report.md"
OUTPUT_DOCX = PROJECT_ROOT / "reports" / "user_experiment_report.docx"

CHINESE_BODY_FONT = "宋体"
CHINESE_HEADING_FONT = "黑体"
ENGLISH_FONT = "Times New Roman"
CODE_FONT = "Consolas"

BODY_SIZE = Pt(12)  # 小四号
HEADING1_SIZE = Pt(14)  # 四号
HEADING2_SIZE = Pt(12)  # 小四号


def set_run_font(run, chinese_font: str, english_font: str, size: Pt, bold: bool = False) -> None:
    """设置 run 的中英文字体。

    Word 中中文字体和西文字体分别存储在不同 OOXML 属性里，因此需要同时设置
    `w:eastAsia` 与普通字体名称，才能让中英文混排时接近截图中的效果。
    """

    run.font.name = english_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), chinese_font)
    run.font.size = size
    run.bold = bold


def set_paragraph_format(paragraph, first_line_indent: bool = True) -> None:
    """设置正文段落格式：小四、首行缩进、1.5 倍行距。"""

    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if first_line_indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)


def add_mixed_text(paragraph, text: str, chinese_font: str, size: Pt, bold: bool = False) -> None:
    """向段落写入中英混排文本。

    Markdown 中的反引号仅表示代码或术语强调，写入 Word 时去掉反引号，并用
    同一正文样式呈现，避免报告里出现多余标记。
    """

    clean_text = text.replace("`", "")
    run = paragraph.add_run(clean_text)
    set_run_font(run, chinese_font=chinese_font, english_font=ENGLISH_FONT, size=size, bold=bold)


def set_cell_shading(cell, fill: str) -> None:
    """设置表格单元格底色。"""

    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, margin_twips: int = 120) -> None:
    """设置表格单元格内边距，避免文字贴边。"""

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def add_table(document: Document, rows: list[list[str]]) -> None:
    """添加 Markdown 表格对应的 Word 表格。"""

    if not rows:
        return

    table = document.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for index, value in enumerate(rows[0]):
        write_cell(header_cells[index], value, is_header=True)

    for row_values in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            write_cell(cells[index], value, is_header=False)

    document.add_paragraph()


def write_cell(cell, text: str, is_header: bool) -> None:
    """写入表格单元格，并设置表头与正文样式。"""

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    if is_header:
        set_cell_shading(cell, "EDEDED")

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text.replace("`", ""))
    set_run_font(
        run,
        chinese_font=CHINESE_HEADING_FONT if is_header else CHINESE_BODY_FONT,
        english_font=ENGLISH_FONT,
        size=Pt(10.5),
        bold=is_header,
    )


def is_table_separator(line: str) -> bool:
    """判断 Markdown 表格分隔行。"""

    return bool(re.fullmatch(r"\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*", line))


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    """从 Markdown 行中解析连续表格。"""

    table_lines: list[str] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        if not is_table_separator(lines[index].strip()):
            cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
            table_lines.append(cells)
        index += 1
    return table_lines, index


def add_heading(document: Document, text: str, level: int) -> None:
    """添加标题，并按截图规则设置标题字体。"""

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_before = Pt(12 if level <= 2 else 6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if level == 1:
        size = HEADING1_SIZE
        chinese_font = CHINESE_HEADING_FONT
        bold = True
    elif level == 2:
        size = HEADING2_SIZE
        chinese_font = CHINESE_HEADING_FONT
        bold = True
    else:
        size = BODY_SIZE
        chinese_font = CHINESE_BODY_FONT
        bold = True

    add_mixed_text(paragraph, text, chinese_font=chinese_font, size=size, bold=bold)


def add_body_paragraph(document: Document, text: str) -> None:
    """添加正文段落。"""

    paragraph = document.add_paragraph()
    set_paragraph_format(paragraph, first_line_indent=True)
    add_mixed_text(paragraph, text, chinese_font=CHINESE_BODY_FONT, size=BODY_SIZE)


def add_list_item(document: Document, text: str) -> None:
    """添加项目符号段落。"""

    paragraph = document.add_paragraph(style=None)
    paragraph.paragraph_format.left_indent = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(-12)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    add_mixed_text(paragraph, f"• {text}", chinese_font=CHINESE_BODY_FONT, size=BODY_SIZE)


def add_code_block(document: Document, code_lines: list[str]) -> None:
    """添加代码块。

    代码块使用等宽字体和浅灰底色，便于用户复制 Windows GPU 训练命令。
    """

    for code_line in code_lines:
        # Markdown 为了在 PowerShell 中换行使用了反引号，Word 报告中直接换行即可，
        # 因此去掉行尾反引号，避免读者误以为它是命令的一部分。
        display_line = code_line.rstrip()
        if display_line.endswith("`"):
            display_line = display_line[:-1].rstrip()

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(display_line)
        set_run_font(run, chinese_font=CODE_FONT, english_font=CODE_FONT, size=Pt(10.5))
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F5F5F5")
        paragraph._p.get_or_add_pPr().append(shading)

    document.add_paragraph()


def configure_document(document: Document) -> None:
    """设置页面、默认样式和页码。"""

    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)

    style = document.styles["Normal"]
    style.font.name = ENGLISH_FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_BODY_FONT)
    style.font.size = BODY_SIZE

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_run = header_paragraph.add_run()
    set_run_font(page_run, chinese_font=CHINESE_BODY_FONT, english_font=ENGLISH_FONT, size=Pt(10.5))
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    page_run._r.append(field_begin)
    page_run._r.append(instr_text)
    page_run._r.append(field_end)


def build_docx() -> None:
    """读取 Markdown 报告并生成 Word 文档。"""

    markdown = SOURCE_MD.read_text(encoding="utf-8")
    lines = markdown.splitlines()

    document = Document()
    configure_document(document)

    index = 0
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.strip()

        if line.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(raw_line)
            index += 1
            continue

        if not line:
            index += 1
            continue

        if line.startswith("|"):
            table_rows, next_index = parse_table(lines, index)
            add_table(document, table_rows)
            index = next_index
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line[level:].strip()
            add_heading(document, text, level)
        elif line.startswith("- "):
            add_list_item(document, line[2:].strip())
        else:
            add_body_paragraph(document, line)

        index += 1

    document.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_docx()
